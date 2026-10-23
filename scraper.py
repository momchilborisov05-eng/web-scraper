"""
Book Store Scraper - collects every book from books.toscrape.com
(a website made for scraping practice) and saves the data to a
SQLite database and a formatted Excel report.

For each book it collects: title, category, price, star rating,
availability and the product URL.

Usage:
    python scraper.py                       # all 50 categories (~1000 books)
    python scraper.py --categories 3        # only the first 3 categories (quick test)
    python scraper.py --db books.db --excel books.xlsx
"""

import argparse
import sqlite3
import sys
import time
from dataclasses import dataclass, astuple
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/"
HEADERS = {"User-Agent": "Mozilla/5.0 (portfolio scraper; learning project)"}
DELAY_SECONDS = 0.5  # be polite to the server
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}


@dataclass
class Book:
    title: str
    category: str
    price_gbp: float
    rating: int
    in_stock: bool
    url: str


class Scraper:
    def __init__(self):
        self.session = requests.Session()
        self.session.headers.update(HEADERS)

    def get_page(self, url: str, retries: int = 3) -> BeautifulSoup:
        """Download a page and parse it. Retries on network errors."""
        for attempt in range(1, retries + 1):
            try:
                response = self.session.get(url, timeout=20)
                response.raise_for_status()
                time.sleep(DELAY_SECONDS)
                return BeautifulSoup(response.content, "html.parser")
            except requests.RequestException as e:
                if attempt == retries:
                    raise
                print(f"  ! {e} - retrying ({attempt}/{retries})")
                time.sleep(2 * attempt)

    def get_categories(self) -> list[tuple[str, str]]:
        """Return (name, url) for every category in the sidebar."""
        soup = self.get_page(BASE_URL)
        links = soup.select("div.side_categories ul li ul li a")
        return [(a.get_text(strip=True), urljoin(BASE_URL, a["href"])) for a in links]

    def scrape_category(self, name: str, url: str) -> list[Book]:
        """Collect all books in one category, following the 'next' pagination links."""
        books = []
        while url:
            soup = self.get_page(url)
            for card in soup.select("article.product_pod"):
                books.append(self.parse_book(card, name, url))
            next_link = soup.select_one("li.next a")
            url = urljoin(url, next_link["href"]) if next_link else None
        return books

    @staticmethod
    def parse_book(card, category: str, page_url: str) -> Book:
        link = card.select_one("h3 a")
        rating_class = card.select_one("p.star-rating")["class"]
        rating = next((RATINGS[c] for c in rating_class if c in RATINGS), 0)
        price_text = card.select_one("p.price_color").get_text(strip=True)
        return Book(
            title=link["title"],
            category=category,
            price_gbp=float(price_text.replace("£", "").replace("Â", "")),
            rating=rating,
            in_stock="in stock" in card.select_one("p.availability").get_text().lower(),
            url=urljoin(page_url, link["href"]),
        )


def save_to_sqlite(books: list[Book], db_path: Path) -> None:
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            """CREATE TABLE IF NOT EXISTS books (
                   id INTEGER PRIMARY KEY AUTOINCREMENT,
                   title TEXT NOT NULL,
                   category TEXT NOT NULL,
                   price_gbp REAL NOT NULL,
                   rating INTEGER NOT NULL,
                   in_stock INTEGER NOT NULL,
                   url TEXT NOT NULL UNIQUE,
                   scraped_at TEXT NOT NULL
               )"""
        )
        now = datetime.now().isoformat(timespec="seconds")
        # UPSERT: re-running the scraper updates prices instead of duplicating rows.
        conn.executemany(
            """INSERT INTO books (title, category, price_gbp, rating, in_stock, url, scraped_at)
               VALUES (?, ?, ?, ?, ?, ?, ?)
               ON CONFLICT(url) DO UPDATE SET
                   price_gbp = excluded.price_gbp,
                   rating = excluded.rating,
                   in_stock = excluded.in_stock,
                   scraped_at = excluded.scraped_at""",
            [(*astuple(b), now) for b in books],
        )


def category_stats(db_path: Path) -> list[tuple]:
    """Aggregate statistics per category, calculated with SQL."""
    with sqlite3.connect(db_path) as conn:
        return conn.execute(
            """SELECT category,
                      COUNT(*)                 AS books,
                      ROUND(AVG(price_gbp), 2) AS avg_price,
                      MIN(price_gbp)           AS min_price,
                      MAX(price_gbp)           AS max_price,
                      ROUND(AVG(rating), 2)    AS avg_rating
               FROM books
               GROUP BY category
               ORDER BY books DESC"""
        ).fetchall()


def write_sheet(ws, headers: list[str], rows: list[tuple]) -> None:
    """Write a table with a bold header, autofilter, frozen header row and fitted columns."""
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, header in enumerate(headers, start=1):
        values = [str(header)] + [str(r[i - 1]) for r in rows]
        ws.column_dimensions[get_column_letter(i)].width = min(max(map(len, values)) + 2, 60)


def save_to_excel(books: list[Book], stats: list[tuple], excel_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Books"
    write_sheet(
        ws,
        ["Title", "Category", "Price (GBP)", "Rating (1-5)", "In stock", "URL"],
        [(b.title, b.category, b.price_gbp, b.rating, "Yes" if b.in_stock else "No", b.url)
         for b in books],
    )
    for cell in ws["C"][1:]:
        cell.number_format = "£#,##0.00"

    write_sheet(
        wb.create_sheet("Summary by category"),
        ["Category", "Books", "Avg price", "Min price", "Max price", "Avg rating"],
        stats,
    )
    wb.save(excel_path)


def main():
    sys.stdout.reconfigure(encoding="utf-8")  # Windows consoles default to a legacy code page
    parser = argparse.ArgumentParser(description="Scrape books.toscrape.com into SQLite and Excel.")
    parser.add_argument("--categories", type=int, help="limit to the first N categories")
    parser.add_argument("--db", type=Path, default=Path("books.db"))
    parser.add_argument("--excel", type=Path, default=Path("books.xlsx"))
    args = parser.parse_args()

    scraper = Scraper()
    start = time.time()
    categories = scraper.get_categories()[: args.categories]
    print(f"Found {len(categories)} categories.\n")

    books: list[Book] = []
    for i, (name, url) in enumerate(categories, start=1):
        category_books = scraper.scrape_category(name, url)
        books.extend(category_books)
        print(f"[{i:>2}/{len(categories)}] {name:<22} {len(category_books):>4} books")

    save_to_sqlite(books, args.db)
    stats = category_stats(args.db)
    save_to_excel(books, stats, args.excel)

    print(f"\nDone: {len(books)} books in {time.time() - start:.0f}s")
    print(f"  SQLite database: {args.db}")
    print(f"  Excel report:    {args.excel}")
    if books:
        cheapest = min(books, key=lambda b: b.price_gbp)
        print(f"  Cheapest book:   {cheapest.title} (£{cheapest.price_gbp:.2f})")


if __name__ == "__main__":
    main()
